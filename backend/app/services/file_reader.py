import io
import logging
import re

from pypdf import PdfReader

logger = logging.getLogger(__name__)

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover
    DOCX_AVAILABLE = False


def extract_text_from_bytes(filename: str, data: bytes) -> str:
    """Extract plain text from an uploaded resume file."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _from_pdf(data)
    if name.endswith(".docx"):
        return _from_docx(data)
    if name.endswith((".txt", ".md", ".csv", ".json")):
        return _from_txt(data)
    raise ValueError(
        "Unsupported file type. Please upload a PDF, DOCX, or TXT file."
    )


def _from_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(pages)
        if not text.strip():
            raise ValueError("The PDF appears to be a scanned image with no extractable text.")
        return text
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Could not read PDF file: {exc}")


def _from_docx(data: bytes) -> str:
    if not DOCX_AVAILABLE:
        raise ValueError("DOCX support is not installed on this server.")
    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)
    except Exception as exc:
        raise ValueError(f"Could not read DOCX file: {exc}")


def _from_txt(data: bytes) -> str:
    for encoding in ("utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode text file.")


def clean_text(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text or "").strip()